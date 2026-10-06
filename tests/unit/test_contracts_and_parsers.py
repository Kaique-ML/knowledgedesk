from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pydantic import ValidationError

from knowledgedesk.generation.contracts import GroundedAnswer
from knowledgedesk.generation.relevance import keep_most_relevant_blocks
from knowledgedesk.ingestion.chunking import chunk_blocks
from knowledgedesk.ingestion.parsers import (
    DocumentLimitExceeded,
    ExtractionBlock,
    ParseLimits,
    parse_document,
)


def test_answer_contract_rejects_unknown_evidence_ids() -> None:
    answer = GroundedAnswer.model_validate(
        {
            "status": "answered",
            "answer_blocks": [{"text": "A regra é 30 dias.", "evidence_ids": ["e1"]}],
        }
    )

    with pytest.raises(ValueError, match="not supplied"):
        answer.validate_evidence_ids({"e2"})


def test_answer_contract_requires_clarification_prompt() -> None:
    with pytest.raises(ValidationError, match="clarification question"):
        GroundedAnswer(status="needs_clarification")


def test_adjacent_generated_facts_are_removed() -> None:
    answer = GroundedAnswer.model_validate(
        {
            "status": "answered",
            "answer_blocks": [
                {"text": "A garantia dura 24 meses.", "evidence_ids": ["warranty"]},
                {"text": "A manutencao ocorre a cada 6 meses.", "evidence_ids": ["maintenance"]},
            ],
        }
    )

    filtered = keep_most_relevant_blocks(
        answer,
        "Qual e o prazo de garantia do AX-120?",
        {
            "warranty": "Garantia AX-120 prazo de 24 meses",
            "maintenance": "Manutencao AX-120 a cada 6 meses",
        },
    )

    assert [block.evidence_ids for block in filtered.answer_blocks] == [["warranty"]]


def test_markdown_preserves_heading_and_line_range(tmp_path: Path) -> None:
    path = tmp_path / "policy.md"
    path.write_text("# Returns\n\nSubmit within 30 days.\nRetain the receipt.\n", encoding="utf-8")

    parsed = parse_document(path)

    assert parsed.format == "markdown"
    assert parsed.sha256
    assert parsed.blocks[0].text == "Submit within 30 days.\nRetain the receipt."
    assert parsed.blocks[0].locator == {
        "heading": "Returns",
        "line_start": 3,
        "line_end": 4,
    }


def test_docx_table_keeps_header_values_and_cell_origin(tmp_path: Path) -> None:
    document = DocxDocument()
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Code"
    table.rows[0].cells[1].text = "Warranty days"
    table.rows[1].cells[0].text = "AX-120"
    table.rows[1].cells[1].text = "30"
    path = tmp_path / "catalog.docx"
    document.save(str(path))

    parsed = parse_document(path)

    assert parsed.blocks == (
        ExtractionBlock(
            "Code: AX-120; Warranty days: 30",
            {"heading": "Document", "table": 1, "row": 2},
            "docx-table",
        ),
    )


def test_xlsx_preserves_headers_and_flags_uncached_formula(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prices"
    sheet.append(["SKU", "Amount"])
    sheet.append(["AX-120", "=1+2"])
    path = tmp_path / "catalog.xlsx"
    workbook.save(path)

    parsed = parse_document(path)

    assert (
        parsed.blocks[0].text == "SKU: AX-120; Amount: [formula present; cached result unavailable]"
    )
    assert parsed.blocks[0].locator == {"sheet": "Prices", "range": "A2:B2"}
    assert "Prices!B2" in parsed.warnings[0]


def test_xlsx_accepts_trailing_empty_cells(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Code", "Quantity", "Optional"])
    sheet.append(["AX-120", 14, None])
    path = tmp_path / "trailing-empty.xlsx"
    workbook.save(path)

    parsed = parse_document(path)

    assert parsed.blocks[0].locator == {"sheet": "Sheet", "range": "A2:C2"}
    assert "AX-120" in parsed.blocks[0].text


def test_chunking_preserves_multi_block_provenance_and_hash() -> None:
    chunks = chunk_blocks(
        (
            ExtractionBlock("Heading and first sentence.", {"page": 1, "block": 1}, "pdf-text"),
            ExtractionBlock("Second sentence.", {"page": 1, "block": 2}, "pdf-text"),
        ),
        max_characters=200,
    )

    assert len(chunks) == 1
    assert chunks[0].content == "Heading and first sentence.\n\nSecond sentence."
    assert chunks[0].locator == {"parts": [{"page": 1, "block": 1}, {"page": 1, "block": 2}]}
    assert len(chunks[0].content_sha256) == 64


def test_upload_limit_is_checked_before_parsing(tmp_path: Path) -> None:
    path = tmp_path / "large.md"
    path.write_text("# heading\n" + "x" * 120, encoding="utf-8")

    with pytest.raises(DocumentLimitExceeded, match="byte limit"):
        parse_document(path, ParseLimits(max_bytes=32))
