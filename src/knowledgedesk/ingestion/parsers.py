from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Literal


class UnsupportedDocument(ValueError):
    pass


class DocumentLimitExceeded(ValueError):
    pass


class OcrUnavailable(RuntimeError):
    pass


SupportedFormat = Literal["pdf", "image", "markdown", "docx", "xlsx"]


@dataclass(frozen=True, slots=True)
class ParseLimits:
    max_bytes: int = 20 * 1024 * 1024
    max_archive_bytes: int = 100 * 1024 * 1024
    max_pixels: int = 40_000_000
    max_pages: int = 250
    max_blocks: int = 20_000


@dataclass(frozen=True, slots=True)
class ExtractionBlock:
    text: str
    locator: dict[str, object]
    extraction_method: str


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    format: SupportedFormat
    sha256: str
    blocks: tuple[ExtractionBlock, ...]
    warnings: tuple[str, ...] = ()


def _check_file(path: Path, limits: ParseLimits) -> bytes:
    if path.stat().st_size > limits.max_bytes:
        raise DocumentLimitExceeded("file exceeds configured upload byte limit")
    return path.read_bytes()


def _check_archive(data: bytes, limits: ParseLimits) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
            if total > limits.max_archive_bytes:
                raise DocumentLimitExceeded("expanded archive exceeds configured limit")
    except zipfile.BadZipFile as error:
        raise UnsupportedDocument("invalid OOXML archive") from error


def detect_format(data: bytes, filename: str, limits: ParseLimits) -> SupportedFormat:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"\x89PNG\r\n\x1a\n") or data.startswith(b"\xff\xd8\xff"):
        return "image"
    if data.startswith(b"PK\x03\x04"):
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile as error:
            raise UnsupportedDocument("invalid OOXML archive") from error
        _check_archive(data, limits)
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        raise UnsupportedDocument("ZIP archive is not a supported DOCX or XLSX document")
    if Path(filename).suffix.lower() in {".md", ".markdown"}:
        try:
            data.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise UnsupportedDocument("Markdown must be valid UTF-8") from error
        return "markdown"
    raise UnsupportedDocument("file content is not a supported PDF, image, Markdown, DOCX, or XLSX")


def parse_document(path: Path, limits: ParseLimits | None = None) -> ParsedDocument:
    limits = limits or ParseLimits()
    data = _check_file(path, limits)
    format_name = detect_format(data, path.name, limits)
    digest = hashlib.sha256(data).hexdigest()
    if format_name == "pdf":
        blocks, warnings = _parse_pdf(data, limits)
    elif format_name == "image":
        blocks, warnings = _parse_image(data, limits)
    elif format_name == "markdown":
        blocks, warnings = _parse_markdown(data)
    elif format_name == "docx":
        blocks, warnings = _parse_docx(data)
    else:
        blocks, warnings = _parse_xlsx(data)
    if len(blocks) > limits.max_blocks:
        raise DocumentLimitExceeded("document exceeds configured extraction block limit")
    if not any(block.text.strip() for block in blocks):
        raise UnsupportedDocument("no readable text was extracted from this document")
    return ParsedDocument(format_name, digest, tuple(blocks), tuple(warnings))


def _ocr(image: object) -> list[ExtractionBlock]:
    try:
        import pytesseract
        from pytesseract import Output
    except ImportError as error:
        raise OcrUnavailable("OCR runtime dependency is not installed") from error
    try:
        result = pytesseract.image_to_data(image, output_type=Output.DICT, config="--psm 3")
    except Exception as error:
        raise OcrUnavailable(
            "Tesseract OCR failed; confirm the OCR binary and language data"
        ) from error

    lines: dict[tuple[int, int, int], list[int]] = {}
    for index, raw in enumerate(result["text"]):
        if str(raw).strip() and int(result["conf"][index]) >= 0:
            key = (
                int(result["block_num"][index]),
                int(result["par_num"][index]),
                int(result["line_num"][index]),
            )
            lines.setdefault(key, []).append(index)

    blocks = []
    for indexes in lines.values():
        text = " ".join(str(result["text"][index]).strip() for index in indexes)
        left = min(int(result["left"][index]) for index in indexes)
        top = min(int(result["top"][index]) for index in indexes)
        right = max(int(result["left"][index]) + int(result["width"][index]) for index in indexes)
        bottom = max(int(result["top"][index]) + int(result["height"][index]) for index in indexes)
        blocks.append(ExtractionBlock(text, {"bbox": [left, top, right, bottom]}, "tesseract-ocr"))
    if not blocks:
        raise UnsupportedDocument("OCR completed but found no readable text")
    return blocks


def _parse_image(data: bytes, limits: ParseLimits) -> tuple[list[ExtractionBlock], list[str]]:
    from PIL import Image, UnidentifiedImageError

    try:
        image = Image.open(io.BytesIO(data))
        image.verify()
        image = Image.open(io.BytesIO(data))
    except (UnidentifiedImageError, OSError) as error:
        raise UnsupportedDocument("invalid or corrupted image") from error
    if image.width * image.height > limits.max_pixels:
        raise DocumentLimitExceeded("image exceeds configured pixel limit")
    return _ocr(image.convert("RGB")), []


def _parse_pdf(data: bytes, limits: ParseLimits) -> tuple[list[ExtractionBlock], list[str]]:
    import fitz
    from PIL import Image

    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as error:
        raise UnsupportedDocument("invalid or corrupted PDF") from error
    if document.is_encrypted:
        raise UnsupportedDocument("encrypted PDFs are not supported")
    if len(document) > limits.max_pages:
        raise DocumentLimitExceeded("PDF exceeds configured page limit")

    blocks = []
    warnings = []
    for page_index, page in enumerate(document):
        page_text = page.get_text("text").strip()
        if len(page_text) >= 20:
            for block_index, raw in enumerate(page.get_text("blocks")):
                text = str(raw[4]).strip()
                if text:
                    blocks.append(
                        ExtractionBlock(
                            text,
                            {
                                "page": page_index + 1,
                                "bbox": [round(value, 2) for value in raw[:4]],
                                "block": block_index + 1,
                            },
                            "pymupdf-text",
                        )
                    )
            continue
        pixels = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        image = Image.frombytes("RGB", (pixels.width, pixels.height), pixels.samples)
        page_blocks = _ocr(image)
        for block_number, block in enumerate(page_blocks, start=1):
            blocks.append(
                ExtractionBlock(
                    block.text,
                    {"page": page_index + 1, "block": block_number, **block.locator},
                    block.extraction_method,
                )
            )
        warnings.append(f"page {page_index + 1} required OCR; layout/table fidelity may be reduced")
    return blocks, warnings


def _parse_markdown(data: bytes) -> tuple[list[ExtractionBlock], list[str]]:
    lines = data.decode("utf-8-sig").splitlines()
    heading = "Document"
    blocks = []
    paragraph: list[str] = []
    start_line = 1

    def flush(end_line: int) -> None:
        if paragraph:
            blocks.append(
                ExtractionBlock(
                    "\n".join(paragraph),
                    {"heading": heading, "line_start": start_line, "line_end": end_line},
                    "markdown-lines",
                )
            )
            paragraph.clear()

    for line_number, line in enumerate(lines, start=1):
        match = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", line)
        if match:
            flush(line_number - 1)
            heading = match.group(1)
            start_line = line_number
            continue
        if not line.strip():
            flush(line_number - 1)
            start_line = line_number + 1
        else:
            if not paragraph:
                start_line = line_number
            paragraph.append(line)
    flush(len(lines))
    return blocks, []


def _parse_docx(data: bytes) -> tuple[list[ExtractionBlock], list[str]]:
    from docx import Document as DocxDocument

    document = DocxDocument(io.BytesIO(data))
    blocks = []
    heading = "Document"
    paragraph_index = 0
    for paragraph in document.paragraphs:
        paragraph_index += 1
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style is not None and paragraph.style.name.startswith("Heading"):
            heading = text
        else:
            blocks.append(
                ExtractionBlock(
                    text,
                    {"heading": heading, "paragraph": paragraph_index},
                    "docx-paragraph",
                )
            )
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        headers = rows[0]
        for row_index, values in enumerate(rows[1:], start=2):
            text = "; ".join(
                f"{header}: {value}"
                for header, value in zip(headers, values, strict=False)
                if header or value
            )
            if text:
                blocks.append(
                    ExtractionBlock(
                        text,
                        {"heading": heading, "table": table_index, "row": row_index},
                        "docx-table",
                    )
                )
    return blocks, []


def _parse_xlsx(data: bytes) -> tuple[list[ExtractionBlock], list[str]]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks = []
    warnings: list[str] = []
    for formula_sheet, value_sheet in zip(formulas.worksheets, cached.worksheets, strict=True):
        rows = formula_sheet.iter_rows()
        value_rows = value_sheet.iter_rows()
        try:
            headers = next(value_rows)
            next(rows)
        except StopIteration:
            continue
        header_values = [
            str(cell.value) if cell.value is not None else cell.coordinate for cell in headers
        ]
        for row_number, (formula_row, value_row) in enumerate(
            zip(rows, value_rows, strict=False), start=2
        ):
            fields = []
            for header, formula_cell, value_cell in zip(
                header_values, formula_row, value_row, strict=False
            ):
                if formula_cell.value is None and value_cell.value is None:
                    continue
                if formula_cell.data_type == "f" and value_cell.value is None:
                    value = "[formula present; cached result unavailable]"
                    warnings.append(
                        f"{formula_sheet.title}!{formula_cell.coordinate} contains a formula "
                        "without a cached result"
                    )
                else:
                    value = str(value_cell.value)
                    if formula_cell.data_type == "f":
                        warnings.append(
                            f"{formula_sheet.title}!{formula_cell.coordinate} uses a cached "
                            "formula value; it was not recalculated"
                        )
                fields.append(f"{header}: {value}")
            if fields:
                last_column = get_column_letter(len(formula_row)) if formula_row else "A"
                cell_range = f"A{row_number}:{last_column}{row_number}"
                blocks.append(
                    ExtractionBlock(
                        "; ".join(fields),
                        {"sheet": formula_sheet.title, "range": cell_range},
                        "xlsx-cached-values",
                    )
                )
    formulas.close()
    cached.close()
    return blocks, sorted(set(warnings))
